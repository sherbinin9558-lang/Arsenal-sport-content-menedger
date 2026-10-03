
import streamlit as st
from PIL import Image, ImageDraw, ImageFont
import json, io, datetime, csv, requests, base64, re, tempfile, os
from pathlib import Path
from max_features import product_search, catalog_metrics, auto_content_bundle, planner_suggestions, knowledge_answer
from ai_seller import ai_sales_reply, sales_followup
from crm_core import create_lead, crm_metrics, load_leads, update_lead, add_lead_interaction, set_customer_profile, STATUSES as CRM_STATUSES
from free_automation import load_orders, create_order, update_order, order_metrics, low_stock, customer_history, content_bundle, seven_day_plan, conversion_metrics
from automation_suite import low_stock_products, stock_info, content_for_product, make_30_day_plan, bulk_update, analytics as automation_analytics, save_uploaded_photo, product_key
from content_manager import WORKFLOW_STATUSES, ensure_workflow, change_status, adapt_content, workflow_metrics, recommendations, report_lines
from growth_engine import ai_summary, attribution_performance, recommendations as growth_recommendations
from saas_core import require_saas_access, render_account_bar, data_load, data_save, data_load_page, data_update_record, data_delete_record, DataConflictError, saas_enabled, tenant_plan, feature_allowed, activate_paid_subscription, can, platform_admin_enabled, platform_admin_snapshot, platform_admin_set_tenant, platform_admin_set_subscription
from webmcp_tools import mount_webmcp_tools
from ui.catalog import render_catalog
from ui.dashboard import render_dashboard
from ui.settings import render_settings

# ==================== КОНФИГУРАЦИЯ ====================
PRODUCTS_FILE = Path("products.json")
CONTENT_PLAN_FILE = Path("content_plan.json")
CARD_SIZE = (1080, 1350)
FONT_REGULAR = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"

CATEGORIES = ["Футболка", "Кроссовки", "Спортивный костюм", "Аксессуары", "Инвентарь", "Другое"]
TONES = ["Официальный", "Дружеский", "Продающий"]
PLATFORMS = ["Instagram", "Telegram", "VK", "Другое"]
CONTENT_TYPES = ["Пост", "Reels", "Stories", "Карусель"]
STATUSES = WORKFLOW_STATUSES
PRIORITIES = ["Обычный", "Высокий", "Срочно"]

TEMPLATES = {
    # Новые варианты.
    "Dark Premium":       {"bg": (9,12,18), "text": (248,249,252), "sec": (148,157,175), "accent": (105,120,255)},
    "Sport Performance":  {"bg": (8,16,27), "text": (248,250,255), "sec": (160,181,207), "accent": (65,160,255)},
    "Editorial Sport":    {"bg": (244,246,249), "text": (20,24,32), "sec": (91,99,113), "accent": (35,71,150)},
    "Black & Electric":   {"bg": (4,5,8), "text": (250,250,252), "sec": (135,140,151), "accent": (168,92,255)},
    # Совместимость со старыми сохранёнными товарами и Reels.
    "Спортивный":         {"bg": (18,25,36), "text": (245,247,250), "sec": (166,181,201), "accent": (105,164,235)},
    "Премиум":            {"bg": (18,20,27), "text": (239,242,247), "sec": (177,184,198), "accent": (137,125,255)},
}

CATEGORY_EMOJI = {
    "Футболка": "👕", "Кроссовки": "👟", "Спортивный костюм": "🩳",
    "Аксессуары": "🧢", "Инвентарь": "🏋️", "Другое": "📦"
}

# ==================== ДАННЫЕ ====================
def load_json(path, default):
    if path.exists():
        try:
            with open(path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return default
    return default

def save_json(path, data):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=4)

def load_products():
    key = "_app_products_cache"
    if key not in st.session_state:
        st.session_state[key] = data_load("products", [])
    return st.session_state[key]

def save_products(p):
    data_save("products", p)
    st.session_state["_app_products_cache"] = p
def add_product(prod):
    p = load_products()
    if not feature_allowed("products", len(p)):
        raise ValueError(f"Лимит каталога тарифа {tenant_plan().upper()} достигнут.")
    p.append(prod); save_products(p)
def _find_record_index(rows, record_id):
    if isinstance(record_id, int):
        return record_id if 0 <= record_id < len(rows) else None
    target = str(record_id or "").strip()
    if not target:
        return None
    for idx, row in enumerate(rows):
        if str(row.get("_saas_record_id", "")) == target:
            return idx
    return None

def update_product(record_id, prod, existing=None):
    if existing is not None and existing.get("_saas_record_id") and existing.get("_saas_updated_at") and saas_enabled():
        merged = dict(existing)
        merged.update(dict(prod or {}))
        merged.pop("_saas_updated_at", None)
        data_update_record("products", existing["_saas_record_id"], merged, existing["_saas_updated_at"])
        st.session_state.pop("_app_products_cache", None)
        return
    p = load_products()
    idx = _find_record_index(p, record_id)
    if idx is not None:
        current = dict(p[idx] or {})
        updated = dict(prod or {})
        stable_id = current.get("_saas_record_id")
        current.update(updated)
        if stable_id:
            current["_saas_record_id"] = stable_id
        else:
            current.pop("_saas_record_id", None)
        p[idx] = current
        save_products(p)

def delete_product(record_id, existing=None):
    if existing is not None and existing.get("_saas_record_id") and existing.get("_saas_updated_at") and saas_enabled():
        data_delete_record("products", existing["_saas_record_id"], existing["_saas_updated_at"])
        st.session_state.pop("_app_products_cache", None)
        return
    p = load_products()
    idx = _find_record_index(p, record_id)
    if idx is not None:
        p.pop(idx)
        save_products(p)

def load_plan():
    key = "_app_plan_cache"
    if key not in st.session_state:
        st.session_state[key] = data_load("content_plan", [])
    return st.session_state[key]

def save_plan(pl):
    data_save("content_plan", pl)
    st.session_state["_app_plan_cache"] = pl

ATTRIBUTION_FILE = Path("content_attribution.json")
def load_attribution(): return data_load("content_attribution", [])
def save_attribution(items): data_save("content_attribution", items)

def attribution_metrics(attribution, leads, orders):
    by_product = {}
    for item in attribution:
        key = str(item.get("product_id") or item.get("product") or "")
        if not key:
            continue
        row = by_product.setdefault(key, {"content": 0, "leads": 0, "orders": 0})
        row["content"] += 1
        row["leads"] += int(item.get("leads", 0) or 0)
        row["orders"] += int(item.get("orders", 0) or 0)
    return by_product

def add_plan(item):
    p = load_plan()
    if not feature_allowed("content", len(p)):
        raise ValueError(f"Лимит контента тарифа {tenant_plan().upper()} достигнут.")
    p.append(item); save_plan(p)
def update_plan(record_id, item):
    p = load_plan()
    idx = _find_record_index(p, record_id)
    if idx is not None:
        existing = dict(p[idx] or {})
        updated = dict(item or {})
        stable_id = existing.get("_saas_record_id")
        existing.update(updated)
        if stable_id:
            existing["_saas_record_id"] = stable_id
        else:
            existing.pop("_saas_record_id", None)
        p[idx] = existing
        save_plan(p)

def delete_plan(record_id):
    p = load_plan()
    idx = _find_record_index(p, record_id)
    if idx is not None:
        p.pop(idx)
        save_plan(p)

def _tenant_asset_root():
    root = Path("tenant_assets") / str(st.session_state.get("saas_tenant_id", "unknown"))
    root.mkdir(parents=True, exist_ok=True)
    return root

def get_logo():
    path = _tenant_asset_root() / "logo.png"
    if path.exists():
        try: return Image.open(path).convert("RGBA")
        except Exception: return None
    return None

def save_logo(f):
    img = Image.open(f).convert("RGBA")
    img.save(_tenant_asset_root() / "logo.png")

# ==================== КАРТОЧКА ====================
def get_font(size, bold=False):
    try:
        return ImageFont.truetype(FONT_BOLD if bold else FONT_REGULAR, size)
    except Exception:
        return ImageFont.load_default()

def fit_font(draw, text, max_w, max_size=80, min_size=38):
    size = max_size
    while size > min_size:
        f = get_font(size, True)
        bbox = draw.textbbox((0, 0), text, font=f)
        if (bbox[2] - bbox[0]) <= max_w:
            return f
        size -= 3
    return get_font(min_size, True)

def draw_centered(draw, text, font, fill, y, w):
    bbox = draw.textbbox((0, 0), text, font=font)
    draw.text(((w - (bbox[2] - bbox[0])) // 2, y), text, font=font, fill=fill)

def _fit_text_lines(draw, text, font, max_w, max_lines=2):
    words = (text or "").split()
    lines = []
    cur = ""
    for word in words:
        test = f"{cur} {word}".strip()
        if draw.textbbox((0, 0), test, font=font)[2] <= max_w:
            cur = test
        else:
            if cur:
                lines.append(cur)
            cur = word
            if len(lines) >= max_lines:
                break
    if cur and len(lines) < max_lines:
        lines.append(cur)
    return lines


from card_generator import generate_card

# ==================== ТЕКСТЫ ====================
def make_hashtags(name, brand, category):
    base = ["#ai_agent_content_manager", "#спорт", "#экипировка", "#новинка"]
    extra = {
        "Футболка": ["#футболка", "#tshirt"],
        "Кроссовки": ["#кроссовки", "#sneakers"],
        "Спортивный костюм": ["#спорткостюм", "#sportswear"],
        "Аксессуары": ["#аксессуары", "#accessories"],
        "Инвентарь": ["#инвентарь", "#gym"],
        "Другое": ["#спорттовары"],
    }
    tags = base + extra.get(category, ["#спорттовары"])
    if brand:
        tags.append(f"#{brand.lower().replace(' ', '')}")
    return " ".join(tags)

def gen_instagram(p, tone):
    n, b = p.get('name',''), p.get('brand','')
    category = p.get('category','Другое')
    desc = p.get('description','').strip()
    specs = p.get('specs','').strip()
    sizes = p.get('sizes','уточняйте')
    color = p.get('color','уточняйте')
    tags = make_hashtags(n, b, category)
    emoji = CATEGORY_EMOJI.get(category, '📦')
    heads = {
        "Официальный": "Новая позиция в AI Agent Content Manager — " + n,
        "Дружеский": "🔥 Забирайте новинку: " + n,
        "Продающий": "🔥 " + n + " — экипировка для тех, кто выбирает по делу!"
    }
    return heads.get(tone, heads["Официальный"]) + f"""

{emoji} {n}
🏷️ Бренд: {b}
🎨 Цвет: {color}
📏 Размеры: {sizes}

{desc}

⚙️ {specs}

📩 Напишите нам в сообщения AI Agent Content Manager — поможем подобрать размер и оформить заказ.

{tags}"""

def gen_telegram(p, tone):
    n, b = p.get('name',''), p.get('brand','')
    category = p.get('category','Другое')
    desc = p.get('description','').strip()
    specs = p.get('specs','').strip()
    sizes = p.get('sizes','уточняйте')
    color = p.get('color','уточняйте')
    emoji = CATEGORY_EMOJI.get(category, '📦')
    heads = {
        "Официальный": "🏆 Новое поступление в AI Agent Content Manager",
        "Дружеский": "🎉 Ребята, смотрите, что приехало!",
        "Продающий": "🔥 НОВИНКА В AI AGENT CONTENT MANAGER"
    }
    return f"""{heads.get(tone, heads["Официальный"])}

{emoji} {n}
Бренд: {b}

{desc}

Характеристики:
• Размеры: {sizes}
• Цвет: {color}
• Особенности: {specs}

📩 Для заказа напишите нам в сообщения. Подскажем наличие и поможем подобрать вариант."""

def gen_vk(p, tone):
    n, b = p.get('name',''), p.get('brand','')
    desc = p.get('description','').strip()
    sizes = p.get('sizes','уточняйте')
    color = p.get('color','уточняйте')
    specs = p.get('specs','').strip()
    return f"""🔥 {n} — {b}

{desc}

📏 Размеры: {sizes}
🎨 Цвет: {color}
⚙️ Характеристики: {specs}

📩 Чтобы заказать товар или уточнить наличие, напишите нам в сообщения сообщества.
{make_hashtags(n, b, p.get('category','Другое'))}"""

# ==================== ПУБЛИКАЦИЯ В TELEGRAM ====================
def publish_to_telegram(image_bytes, caption):
    try:
        token = st.secrets["TELEGRAM_TOKEN"]
        channel = st.secrets["TELEGRAM_CHANNEL"]
        if not channel.startswith("@"):
            channel = "@" + channel
        api_url = f"https://api.telegram.org/bot{token}/sendPhoto"
        files = {"photo": ("card.png", image_bytes, "image/png")}
        data = {"chat_id": channel, "caption": caption[:1024]}
        resp = requests.post(api_url, files=files, data=data, timeout=60)
        if resp.status_code == 200:
            return True, "Пост успешно опубликован!"
        else:
            return False, f"Ошибка {resp.status_code}: {resp.text}"
    except KeyError as e:
        return False, f"Не найден секрет: {e}."
    except Exception as e:
        return False, f"Ошибка публикации: {e}"

# ==================== ПУБЛИКАЦИЯ REELS ====================
def _telegram_chat_id(raw_channel):
    value = str(raw_channel or "").strip()
    if value.startswith("https://t.me/"):
        value = value.rstrip("/").split("/")[-1]
    if value.startswith("t.me/"):
        value = value.rstrip("/").split("/")[-1]
    if value.startswith("-100") or value.lstrip("-").isdigit():
        return value
    return value if value.startswith("@") else "@" + value

def _telegram_api(token, method):
    return f"https://api.telegram.org/bot{token}/{method}"

def publish_reel_to_telegram(video_bytes, caption):
    try:
        token = str(st.secrets["TELEGRAM_TOKEN"]).strip()
        channel = _telegram_chat_id(st.secrets["TELEGRAM_CHANNEL"])
        if not token:
            return False, "TELEGRAM_TOKEN пустой."
        if not channel:
            return False, "TELEGRAM_CHANNEL пустой."

        size_mb = len(video_bytes) / (1024 * 1024)
        if size_mb > 50:
            return False, f"Видео весит {size_mb:.1f} МБ — больше лимита Telegram Bot API 50 МБ."

        caption = (caption or "").strip()[:1024]
        api_url = _telegram_api(token, "sendVideo")

        files = {
            "video": ("ai_agent_content_manager_reel.mp4", video_bytes, "video/mp4")
        }
        data = {
            "chat_id": channel,
            "caption": caption,
            "supports_streaming": "true",
        }
        resp = requests.post(
            api_url,
            files=files,
            data=data,
            timeout=(20, 300),
        )

        try:
            result = resp.json()
        except Exception:
            result = {}

        if resp.ok and result.get("ok") is True:
            return True, f"Reels опубликован в Telegram ({channel}). Размер: {size_mb:.1f} МБ."

        description = result.get("description") or resp.text or "неизвестная ошибка"
        lower = description.lower()

        if "chat not found" in lower:
            return False, "Telegram не нашёл канал. Проверь TELEGRAM_CHANNEL: @username канала или числовой ID вида -100xxxxxxxxxx."
        if "not enough rights" in lower or "forbidden" in lower:
            return False, "Бот найден, но Telegram запретил публикацию. Бот должен быть администратором канала с правом публикации сообщений."
        if "file is too big" in lower:
            return False, f"Telegram отклонил видео как слишком большое ({size_mb:.1f} МБ)."
        if "wrong file identifier/http url specified" in lower:
            return False, f"Telegram не принял MP4. Размер {size_mb:.1f} МБ. Ответ API: {description}"

        # Резервный канал: если sendVideo не прошёл из-за формата/обработки,
        # пробуем загрузить тот же MP4 как документ. Это позволяет не терять файл
        # и одновременно показывает, что проблема именно в обработке video Telegram.
        try:
            doc_resp = requests.post(
                _telegram_api(token, "sendDocument"),
                files={"document": ("ai_agent_content_manager_reel.mp4", video_bytes, "video/mp4")},
                data={
                    "chat_id": channel,
                    "caption": caption,
                },
                timeout=(20, 300),
            )
            try:
                doc_result = doc_resp.json()
            except Exception:
                doc_result = {}

            if doc_resp.ok and doc_result.get("ok") is True:
                return True, (
                    f"Telegram не принял файл как видео, поэтому отправил его как MP4-документ. "
                    f"Размер: {size_mb:.1f} МБ. Причина sendVideo: {description}"
                )
        except requests.RequestException:
            pass

        return False, f"Telegram API не принял Reels: {description} (HTTP {resp.status_code}, {size_mb:.1f} МБ)."

    except KeyError as e:
        return False, f"Не найден секрет: {e}. Нужны TELEGRAM_TOKEN и TELEGRAM_CHANNEL."
    except requests.RequestException as e:
        return False, f"Сеть/Telegram недоступны: {e}"
    except Exception as e:
        return False, f"Ошибка Telegram: {e}"

def publish_reel_to_vk(video_bytes, caption):
    try:
        token = st.secrets["VK_ACCESS_TOKEN"]
        owner_id = int(st.secrets["VK_OWNER_ID"])
        api_version = str(st.secrets.get("VK_API_VERSION", "5.199"))

        save_url = "https://api.vk.com/method/video.save"
        params = {
            "access_token": token,
            "v": api_version,
            "name": "AI Agent Content Manager Reels",
            "description": caption[:4096],
            "is_private": 0,
            "wallpost": 1,
            "owner_id": owner_id,
        }
        resp = requests.post(save_url, data=params, timeout=60)
        data = resp.json()

        if "error" in data:
            return False, f"Ошибка VK: {data['error'].get('error_msg', data['error'])}"

        video = data.get("response", {})
        upload_url = video.get("upload_url")
        if not upload_url:
            return False, "VK не вернул upload_url."

        upload = requests.post(
            upload_url,
            files={"video_file": ("reel.mp4", video_bytes, "video/mp4")},
            timeout=180,
        )
        upload_data = upload.json()
        if "error" in upload_data:
            return False, f"Ошибка загрузки VK: {upload_data['error']}"

        return True, "Reels отправлен в VK!"
    except KeyError as e:
        return False, f"Не найден секрет: {e}."
    except Exception as e:
        return False, f"Ошибка VK: {e}"


# ==================== OWNER CONSOLE ====================
def render_platform_admin():
    st.markdown(
        '<div class="dashboard-hero">'
        '<div class="dashboard-hero-kicker">PLATFORM OWNER</div>'
        '<div class="dashboard-hero-title">Админ-панель платформы</div>'
        '<div class="dashboard-hero-text">Управление магазинами, тарифами, подписками и аккаунтами AI Agent Content Manager.</div>'
        '</div>',
        unsafe_allow_html=True,
    )
    try:
        tenants, subs, members, users = platform_admin_snapshot()
    except Exception as e:
        st.error(f"Не удалось загрузить данные админ-панели: {e}")
        return

    sub_by_tenant = {str(x.get("tenant_id")): x for x in subs}
    owner_by_tenant = {}
    for m in members:
        if m.get("role") == "owner" and str(m.get("tenant_id")) not in owner_by_tenant:
            owner_by_tenant[str(m.get("tenant_id"))] = m.get("user_id")
    email_by_user = {str(u.get("id")): u.get("email","") for u in users}

    active = sum(1 for t in tenants if str(t.get("status","")).lower() == "active")
    trial = sum(1 for t in tenants if str(t.get("plan","")).lower() == "trial")
    paid = sum(1 for t in tenants if str(t.get("plan","")).lower() in ("starter","pro","business"))

    k1,k2,k3,k4 = st.columns(4)
    k1.metric("Магазины", len(tenants))
    k2.metric("Активные", active)
    k3.metric("Trial", trial)
    k4.metric("Платные", paid)

    st.markdown("### Магазины")
    rows = []
    for t in tenants:
        tid = str(t.get("id"))
        sub = sub_by_tenant.get(tid,{})
        owner_id = str(owner_by_tenant.get(tid) or "")
        rows.append({
            "Магазин": t.get("name","—"),
            "Владелец": email_by_user.get(owner_id,"—"),
            "Тариф": str(t.get("plan","trial")).upper(),
            "Статус": t.get("status","—"),
            "Подписка": sub.get("status","—"),
            "Создан": str(t.get("created_at",""))[:10],
        })
    st.dataframe(rows, use_container_width=True, hide_index=True)

    if not tenants:
        st.info("Пока нет зарегистрированных магазинов.")
        return

    st.markdown("### Управление магазином")
    labels = [f'{t.get("name","Магазин")} · {str(t.get("plan","trial")).upper()}' for t in tenants]
    selected = st.selectbox("Магазин", labels, key="platform_admin_tenant")
    t = tenants[labels.index(selected)]
    tid = str(t.get("id"))
    sub = sub_by_tenant.get(tid,{})
    current_plan = str(t.get("plan","trial")).lower()
    current_status = str(t.get("status","active")).lower()
    current_sub_status = str(sub.get("status","trialing")).lower()
    current_renew = bool(sub.get("auto_renew",False))

    a,b = st.columns(2)
    with a:
        new_plan = st.selectbox("Тариф", ["trial","starter","pro","business"],
                                index=["trial","starter","pro","business"].index(current_plan) if current_plan in ["trial","starter","pro","business"] else 0,
                                key="platform_admin_plan")
        new_status = st.selectbox("Статус магазина", ["active","suspended","cancelled"],
                                  index=["active","suspended","cancelled"].index(current_status) if current_status in ["active","suspended","cancelled"] else 0,
                                  key="platform_admin_status")
    with b:
        new_sub_status = st.selectbox("Статус подписки", ["trialing","active","past_due","canceled"],
                                      index=["trialing","active","past_due","canceled"].index(current_sub_status) if current_sub_status in ["trialing","active","past_due","canceled"] else 0,
                                      key="platform_admin_sub_status")
        new_renew = st.toggle("Автопродление", value=current_renew, key="platform_admin_renew")

    if st.button("Сохранить изменения", type="primary", use_container_width=True, key="platform_admin_save"):
        try:
            platform_admin_set_tenant(tid, plan=new_plan, status=new_status)
            platform_admin_set_subscription(tid, plan=new_plan, status=new_sub_status, auto_renew=new_renew)
            st.success("Изменения сохранены.")
            st.rerun()
        except Exception as e:
            st.error(f"Не удалось сохранить: {e}")

    st.markdown("---")
    st.markdown("### Настройки администратора")
    st.caption("Системные настройки платформы. Чувствительные параметры не хранятся и не редактируются через клиентский интерфейс.")
    s1, s2, s3 = st.columns(3)
    s1.metric("Доступ", "Разрешён")
    s2.metric("Режим", "Platform Owner")
    s3.metric("Управление", "Системное")
    st.write("**Администратор платформы:** определяется серверной конфигурацией.")
    st.write("**Безопасность:** секретные ключи и доступ к сервисной роли остаются вне интерфейса.")
    st.caption("Изменение системных секретов выполняется только в настройках окружения. Это предотвращает сохранение чувствительных данных в базе и в клиентском коде.")
    if st.button("↻ Обновить данные админ-панели", key="platform_admin_refresh", use_container_width=True):
        st.rerun()

    st.markdown("---")
    st.markdown("### Аккаунты")
    account_rows = []
    for u in users:
        uid = str(u.get("id"))
        tenant_for_user = next((x for x in members if str(x.get("user_id")) == uid), {})
        account_rows.append({
            "Email": u.get("email","—"),
            "Роль": tenant_for_user.get("role","—"),
            "Подтверждён": "Да" if u.get("email_confirmed_at") else "Нет",
            "Создан": str(u.get("created_at",""))[:10],
        })
    st.dataframe(account_rows, use_container_width=True, hide_index=True)
    st.caption("Вкладка доступна только владельцу платформы из SAAS_ADMIN_EMAIL. Service-role ключ нужен для загрузки и управления данными.")


# ==================== MAX ====================
from ui.max import render_max


# ==================== ИНТЕРФЕЙС ====================
st.set_page_config(page_title="AI Agent Content Manager", page_icon="⚡", layout="wide", initial_sidebar_state="expanded")

st.markdown("""
<style>
button[key="sidebar_max"]{
background:linear-gradient(135deg,#b8ff00,#7cff00)!important;
color:#101500!important;
border:1px solid #d7ff72!important;
box-shadow:0 0 0 1px rgba(184,255,0,.55),0 0 18px rgba(184,255,0,.42)!important;
font-weight:900!important;
}
</style>
""", unsafe_allow_html=True)

if not require_saas_access():
    st.stop()

mount_webmcp_tools()

# Проверка результата оплаты после возврата с ЮKassa.
try:
    payment_id=st.query_params.get("payment_id")
    checkout_id=st.query_params.get("checkout_id")
    billing_return=st.query_params.get("billing")
    if billing_return == "tbank_return" and not st.session_state.get("tbank_verified"):
        from tbank_billing import get_state
        pid=st.session_state.get("tbank_payment_id")
        if pid:
            payment=get_state(pid)
            status=payment.get("Status")
            plan=st.session_state.get("tbank_plan")
            if status in ("CONFIRMED","AUTHORIZED") and plan in ("starter","pro","business"):
                activate_paid_subscription(plan,str(pid),None,provider="tbank",tenant=st.session_state.get("saas_tenant_id"))
                st.session_state["tbank_verified"]=True
                st.success(f"Оплата Т‑Банка подтверждена. Тариф {plan.upper()} активирован.")
            elif status:
                st.info(f"Статус платежа Т‑Банк: {status}.")
    elif (payment_id or checkout_id) and billing_return == "return" and not st.session_state.get("billing_verified"):
        from billing import get_checkout_by_order, get_payment
        if checkout_id:
            checkout=get_checkout_by_order(str(checkout_id))
            if not checkout or not checkout.get("provider_payment_id"):
                raise RuntimeError("Checkout-сессия не найдена или не содержит ID платежа.")
            payment_id=str(checkout["provider_payment_id"])
        payment=get_payment(payment_id)
        if payment.get("status") == "succeeded" and payment.get("paid"):
            plan=str((payment.get("metadata") or {}).get("plan") or st.session_state.get("billing_plan") or "").lower()
            tenant=str((payment.get("metadata") or {}).get("tenant_id") or st.session_state.get("saas_tenant_id") or "")
            if plan in ("starter","pro","business") and tenant == str(st.session_state.get("saas_tenant_id")):
                method_id=(payment.get("payment_method") or {}).get("id")
                activate_paid_subscription(plan,payment_id,method_id,provider="yookassa",tenant=tenant)
                st.session_state["billing_verified"]=True
                st.success(f"Оплата подтверждена. Тариф {plan.upper()} активирован.")
        elif payment.get("status") == "canceled":
            st.warning("Платёж отменён.")
except Exception as e:
    st.warning(f"Статус оплаты пока не подтверждён: {e}")


with st.sidebar:
    st.markdown(
        '<div class="sidebar-brand">'
        '<div class="sidebar-brand-kicker">AI AGENT</div>'
        '<div class="sidebar-brand-title">Content Manager</div>'
        '<div class="sidebar-brand-line"></div>'
        '<div class="sidebar-brand-version">CONTENT STUDIO · v2.7 MAX</div>'
        '</div>',
        unsafe_allow_html=True,
    )
    dark_mode = st.toggle("🌙 Тёмная тема", value=False, key="theme_toggle")
    st.markdown("---")
    st.markdown(
        '<div class="sidebar-max"><div class="sidebar-max-kicker">AI AGENT CONTENT MANAGER</div>'
        '<div class="sidebar-max-title">⚡ MAX</div>'
        '<div class="sidebar-max-text">Центр управления магазином</div></div>',
        unsafe_allow_html=True,
    )
    if st.button("⚡ Открыть MAX", use_container_width=True, type="primary", key="sidebar_max"):
        st.session_state["open_max"] = True
        render_max()
    st.markdown("---")
    try:
        tg_ch = st.secrets.get("TELEGRAM_CHANNEL", None)
        if tg_ch:
            st.success(f"📡 Telegram: @{tg_ch}")
        else:
            st.info("📡 Telegram не настроен")
    except Exception:
        st.info("📡 Telegram не настроен")
    st.markdown("---")
    st.caption(f"Сегодня · {datetime.date.today().strftime('%d.%m.%Y')}")

if dark_mode:
    bg_css = "body {background:#0b0e13;}"
    theme_css = """
    .stApp {background:#0b0e13!important;color:#f5f7fa!important;}
    section[data-testid="stSidebar"] {background:#10131a!important;border-right:1px solid #252b38!important;}section[data-testid="stSidebar"] .sidebar-brand-title{color:#f5f7fa!important;} section[data-testid="stSidebar"] .sidebar-brand-version{color:#8d95a5!important;}
    div[data-baseweb="tab-list"] {background:#11151d!important;border-color:#252b38!important;}
    [data-testid="stFileUploader"] {background:#11151d!important;border-color:#394152!important;}
    """
else:
    bg_css = "body {background:#eef1f5;}"
    theme_css = """
    .stApp {background:#f6f7fb!important;color:#171a21!important;}
    section[data-testid="stSidebar"] {background:#ffffff!important;border-right:1px solid #e4e7ec!important;}
    div[data-baseweb="tab-list"] {background:#ffffff!important;border-color:#e1e5eb!important;}
    [data-testid="stFileUploader"] {background:#f5f7fa!important;border-color:#b9c1ce!important;}
    .sidebar-max {background:linear-gradient(135deg,#f5f2ff,#e9e5ff);border-color:#c9c1ff;}.sidebar-brand-kicker{color:#5b4bd6;} .sidebar-brand-title{color:#21164d;} .sidebar-brand-version{color:#7b8190;}
    .sidebar-max-kicker {color:#5b4bd6;}
    .sidebar-max-title {color:#21164d;}
    .sidebar-max-text {color:#5f6472;}
    .stApp, .stApp * {color:#17151c;}
    .stApp [data-testid="stMarkdownContainer"] p, .stApp [data-testid="stMarkdownContainer"] li, .stApp [data-testid="stMarkdownContainer"] span, .stApp label, .stApp [data-testid="stCaptionContainer"] {color:#17151c !important;}
    .stApp h1, .stApp h2, .stApp h3, .stApp h4, .stApp h5, .stApp h6 {color:#17151c !important;}
    .stApp input, .stApp textarea, .stApp [data-baseweb="select"] * {color:#17151c !important;}
    .stApp input::placeholder, .stApp textarea::placeholder {color:#667080 !important;opacity:1 !important;}
    .stApp button:not([data-baseweb="tab"]) {color:#17151c !important;}
    .stApp button[kind="primary"] {color:#ffffff !important;}
    button[data-baseweb="tab"] {color:#17151c !important;background:#ffffff !important;border-color:#dfe3ea !important;}
    button[data-baseweb="tab"] * {color:#17151c !important;}
    button[data-baseweb="tab"][aria-selected="true"] {color:#32165f !important;background:linear-gradient(135deg,#f1eaff,#e8ddff) !important;border-color:#bda6e8 !important;box-shadow:0 0 0 1px rgba(72,35,125,.16),0 0 11px rgba(83,43,145,.18) !important;}
    button[data-baseweb="tab"][aria-selected="true"] * {color:#32165f !important;}
    .section-kicker {color:#4b247c !important;text-shadow:0 0 7px rgba(75,36,124,.22) !important;}
    .pill {color:#32165f !important;}
    /* Light theme: transparent controls/cards so black text stays readable. */
    .stApp .stButton > button:not([kind="primary"]) {background:transparent !important;color:#17151c !important;border-color:#cfd4dd !important;box-shadow:none !important;}
    .stApp .stButton > button:not([kind="primary"]):hover {background:rgba(91,70,214,.05) !important;color:#17151c !important;border-color:#8b7bc7 !important;box-shadow:0 0 0 1px rgba(75,36,124,.12) !important;}
    div[data-baseweb="tab-list"] {background:transparent !important;box-shadow:none !important;border-color:#dfe3ea !important;}
    button[data-baseweb="tab"] {background:transparent !important;color:#17151c !important;border-color:transparent !important;box-shadow:none !important;}
    button[data-baseweb="tab"] * {color:#17151c !important;}
    button[data-baseweb="tab"][aria-selected="true"] {background:rgba(232,221,255,.55) !important;color:#32165f !important;border-color:#bda6e8 !important;box-shadow:0 0 0 1px rgba(72,35,125,.12),0 0 10px rgba(83,43,145,.12) !important;}
    button[data-baseweb="tab"][aria-selected="true"] * {color:#32165f !important;}
    .stat-box {background:transparent !important;color:#17151c !important;border:1px solid #dfe3ea !important;box-shadow:none !important;}
    .stat-number {color:#17151c !important;}
    .stat-label {color:#596174 !important;}
    [data-testid="stAlert"], div[data-baseweb="notification"] {background:transparent !important;color:#17151c !important;border-color:#dfe3ea !important;}
    [data-testid="stAlert"] *, div[data-baseweb="notification"] * {color:#17151c !important;}

    """

st.markdown(
    "<style>\n" + bg_css + theme_css + """
/* ===== Clean customer UI ===== */
/* 1) Hide Streamlit/GitHub/platform chrome. */
#MainMenu,[data-testid="stDecoration"],[data-testid="stStatusWidget"],[data-testid="stAppDeployButton"],footer,[data-testid="stBottom"],[data-testid="stBottomBlockContainer"],[data-testid="stBottomBlock"]{display:none!important;visibility:hidden!important;height:0!important;min-height:0!important;pointer-events:none!important;}
/* Platform chrome stays hidden so customers see only the product UI. */
/* 2) Native sidebar: stable width on desktop, compact drawer on touch devices. */
section[data-testid="stSidebar"]{display:block!important;visibility:visible!important;opacity:1!important;background:#fff!important;border-right:1px solid #e4e7ec!important;}
@media (min-width:769px){section[data-testid="stSidebar"]{width:280px!important;min-width:280px!important;max-width:280px!important;}section[data-testid="stSidebar"]>div:first-child{width:280px!important;max-width:280px!important;}}
/* 3) Sidebar MAX. */
button[key="sidebar_max"]{background:linear-gradient(135deg,#b8ff00,#7cff00)!important;color:#101500!important;border:1px solid #d7ff72!important;box-shadow:0 0 0 1px rgba(184,255,0,.55),0 0 18px rgba(184,255,0,.42)!important;font-weight:900!important;letter-spacing:.01em!important;} button[key="sidebar_max"]:hover{transform:translateY(-1px)!important;box-shadow:0 0 0 1px rgba(184,255,0,.75),0 0 24px rgba(184,255,0,.58)!important;}
/* 4) Mobile/tablet MAX: normal document flow; never fixed/sticky. */
.mobile-max-launcher{display:none!important;}
@media (max-width:1100px){.mobile-max-launcher{display:block!important;width:100%!important;margin:0 0 10px!important;}.mobile-max-launcher button{width:100%!important;min-height:46px!important;border-radius:12px!important;background:linear-gradient(135deg,#7cff00,#b8ff00)!important;color:#101500!important;border:1px solid #d7ff72!important;box-shadow:0 0 0 1px rgba(184,255,0,.55),0 0 14px rgba(184,255,0,.45)!important;font-weight:900!important;}}
/* Main layout. */
.block-container{width:100%!important;max-width:1500px!important;margin:0 auto!important;padding-left:clamp(.75rem,2.5vw,2.5rem)!important;padding-right:clamp(.75rem,2.5vw,2.5rem)!important;}
.main-title{font-size:2.05rem;font-weight:900;color:#24124f;text-align:left;margin:18px 0 12px;letter-spacing:.01em;text-shadow:0 0 10px rgba(91,70,214,.18);}
.mobile-nav-hint{display:none;color:#7a8494;font-size:.74rem;margin:4px 0 8px;}html,body{overflow-x:hidden!important;}
/* MAX dialog. */
div[data-testid="stDialog"]{display:flex!important;visibility:visible!important;opacity:1!important;position:fixed!important;inset:0!important;z-index:2147483647!important;pointer-events:auto!important;}
div[data-testid="stDialog"]>div,div[data-testid="stDialog"] [role="dialog"],div[role="dialog"]{visibility:visible!important;opacity:1!important;pointer-events:auto!important;}
div[data-testid="stDialog"] [role="dialog"]{display:block!important;position:relative!important;z-index:2147483647!important;background:linear-gradient(180deg,#ffffff 0%,#f8f9fc 100%)!important;color:#17151c!important;border:1px solid #d9ddea!important;border-radius:22px!important;box-shadow:0 28px 90px rgba(12,16,28,.34),0 0 0 1px rgba(91,92,226,.06)!important;max-height:90vh!important;overflow:auto!important;} div[data-testid="stDialog"] [role="dialog"] *{visibility:visible!important;}
[data-testid="stDialog"] [data-testid="stExpander"] button,div[role="dialog"] [data-testid="stExpander"] button{color:#b8ff00!important;opacity:1!important;}
[data-testid="stDialog"] [data-testid="stExpander"] button svg,[data-testid="stDialog"] [data-testid="stExpander"] button svg *,div[role="dialog"] [data-testid="stExpander"] button svg,div[role="dialog"] [data-testid="stExpander"] button svg *{color:#b8ff00!important;stroke:#b8ff00!important;fill:none!important;opacity:1!important;stroke-width:3px!important;filter:drop-shadow(0 0 6px rgba(184,255,0,.9))!important;}
.max-header{display:block!important;width:100%!important;padding:6px 52px 14px 0!important;border-bottom:1px solid #e8ebf1!important;margin-bottom:8px!important;}.max-header-title{font-size:clamp(1.25rem,2vw,1.55rem)!important;line-height:1.2!important;font-weight:950!important;letter-spacing:-.02em!important;color:#21164d!important;}.max-header-title span{color:#7cff00!important;text-shadow:0 0 8px rgba(184,255,0,.55)!important;}.max-header-subtitle{margin-top:4px!important;font-size:.82rem!important;color:#687182!important;}
/* Existing visual styles retained. */
:root{--accent:#5b5ce2;--accent2:#4f46c5;--bg:#f6f7fb;--surface:#fff;--border:#e4e7ec;--text:#171a21;--muted:#687182;}.stApp{background:#f6f7fb;}
.sidebar-ai-agent{font-size:1.02rem;font-weight:900;letter-spacing:.015em;color:#b8a6ff;text-shadow:0 0 7px rgba(169,140,255,.65);margin:0 0 7px;text-align:left;line-height:1.2;}.dashboard-hero{padding:26px 28px;border:1px solid #e0e4ee;border-radius:22px;background:radial-gradient(circle at 92% 10%,rgba(124,255,0,.18),transparent 28%),linear-gradient(135deg,#ffffff 0%,#f3f1ff 100%);box-shadow:0 14px 36px rgba(30,35,55,.08);margin-bottom:16px;}.dashboard-hero-kicker{font-size:.66rem;letter-spacing:.16em;font-weight:900;color:#5b5ce2;}.dashboard-hero-title{font-size:clamp(1.65rem,3vw,2.45rem);font-weight:950;letter-spacing:-.04em;color:#171a21;margin-top:7px;line-height:1.05;}.dashboard-hero-text{max-width:760px;margin-top:9px;font-size:.96rem;line-height:1.5;color:#596174;}.first-run-card{padding:20px 22px;border:1px solid #cfc6ff;border-radius:20px;background:linear-gradient(135deg,#faf8ff,#f1edff);margin:10px 0 14px;}.first-run-kicker{font-size:.65rem;letter-spacing:.14em;font-weight:900;color:#6d4aff;}.first-run-title{font-size:1.35rem;font-weight:900;color:#21164d;margin-top:5px;}.first-run-text{color:#596174;margin-top:5px;}.stAlert{border-radius:14px!important;}.empty-state{padding:24px;border:1px dashed #cdd3df;border-radius:18px;text-align:center;background:rgba(255,255,255,.55);color:#687182;}.max-ai-plan{padding:16px 18px;border:1px solid #cfc6ff;border-radius:18px;background:linear-gradient(135deg,#fbfaff,#f1edff);margin:8px 0 14px;}.max-ai-plan-title{font-weight:900;color:#32165f;font-size:1.05rem;}.max-ai-plan-item{margin-top:8px;color:#4e5361;font-size:.9rem;line-height:1.4;}.sidebar-brand{padding:4px 2px 12px;margin:0 0 4px;}.sidebar-brand-kicker{font-size:.64rem;letter-spacing:.16em;font-weight:900;color:#8f80ff;line-height:1.1;}.sidebar-brand-title{font-size:1.42rem;font-weight:950;letter-spacing:-.035em;color:#21164d;line-height:1.05;margin-top:4px;}.sidebar-brand-line{width:42px;height:3px;border-radius:99px;background:linear-gradient(90deg,#b8ff00,#7cff00);box-shadow:0 0 10px rgba(184,255,0,.45);margin-top:10px;}.sidebar-brand-version{font-size:.58rem;letter-spacing:.09em;font-weight:800;color:#7b8190;margin-top:9px;}.sidebar-max{padding:12px 13px 10px;border:1px solid #34304f;border-radius:16px;background:linear-gradient(145deg,#151a24,#211f35);margin:3px 0 9px;box-shadow:0 10px 28px rgba(12,15,25,.16);}.sidebar-max-kicker{font-size:.6rem;letter-spacing:.13em;font-weight:850;color:#a69cff;}.sidebar-max-title{font-size:1.22rem;font-weight:950;color:#fff;margin-top:3px;letter-spacing:-.01em;}.sidebar-max-text{font-size:.72rem;color:#aeb5c2;margin-top:3px;}
.stButton>button{border-radius:12px;border:1px solid #303747;background:#171c26;color:#f5f7fa;font-weight:700;transition:transform .16s ease,box-shadow .16s ease,border-color .16s ease,background .16s ease;}.stButton>button:hover{transform:translateY(-1px);border-color:#5b5ce2;box-shadow:0 8px 20px rgba(30,35,55,.14);}button[kind="primary"]{background:linear-gradient(135deg,#756be8,#897dff)!important;border:0!important;color:#fff!important;box-shadow:0 8px 22px rgba(91,92,226,.22)!important;}button[kind="primary"]:hover{box-shadow:0 11px 28px rgba(91,92,226,.30)!important;}
div[data-baseweb="tab-list"]{gap:5px;background:transparent!important;padding:5px;border-radius:14px;border:1px solid transparent!important;display:flex!important;visibility:visible!important;opacity:1!important;}button[data-baseweb="tab"]{border-radius:10px;color:#252b3a!important;font-weight:800;visibility:visible!important;opacity:1!important;min-height:40px;background:transparent!important;border-color:transparent!important;}button[data-baseweb="tab"] *{color:#252b3a!important;opacity:1!important;}button[data-baseweb="tab"][aria-selected="true"]{background:transparent!important;color:#4c1d95!important;border-color:transparent!important;}button[data-baseweb="tab"][aria-selected="true"] *{color:#4c1d95!important;opacity:1!important;}
[data-testid="stHorizontalBlock"]{width:100%!important;}[data-testid="stTextInput"],[data-testid="stTextArea"],[data-testid="stSelectbox"],[data-testid="stNumberInput"],[data-testid="stDateInput"],[data-testid="stFileUploader"]{width:100%!important;}.stat-box{background:linear-gradient(135deg,#171c26 0%,#211f35 100%);border:1px solid #30364a;color:#f5f7fa;padding:20px;border-radius:16px;text-align:center;margin:5px;box-shadow:0 8px 24px rgba(0,0,0,.12);}.stat-number{font-size:2.35rem;font-weight:900;margin:0;color:#f5f7fa;}.stat-label{font-size:.9rem;color:#aeb5c2;margin:0;}
.section-kicker{color:#6d4aff!important;text-shadow:0 0 7px rgba(109,74,255,.20);}.section-title{color:#21164d!important;}.stMarkdown h1,.stMarkdown h2,.stMarkdown h3{color:#21164d!important;}[data-testid="stText"],[data-testid="stCaptionContainer"]{color:#596174!important;}
@media (min-width:769px){section[data-testid="stSidebar"]{display:block!important;visibility:visible!important;opacity:1!important;position:relative!important;transform:none!important;width:280px!important;min-width:280px!important;max-width:280px!important;left:0!important;}section[data-testid="stSidebar"]>div:first-child{width:280px!important;max-width:280px!important;}[data-testid="stAppViewContainer"] .main{width:calc(100vw - 280px)!important;max-width:calc(100vw - 280px)!important;}[data-testid="stAppViewContainer"] .main .block-container{width:100%!important;max-width:none!important;margin:0!important;}[data-testid="stSidebarCollapseButton"]{display:block!important;visibility:visible!important;}}
@media (max-width:768px){.dashboard-hero{padding:20px 18px;border-radius:18px;margin-bottom:12px;}.dashboard-hero-title{font-size:1.55rem;}.dashboard-hero-text{font-size:.86rem;}.first-run-card{padding:17px 16px;border-radius:17px;}.first-run-title{font-size:1.12rem;}.max-ai-plan{padding:14px 15px;border-radius:16px;}.max-ai-plan-item{font-size:.82rem;}.mobile-nav-hint{display:block;}.main-title{font-size:1.35rem!important;line-height:1.15!important;margin:8px 0 7px!important;}.block-container{padding:.65rem .7rem 1.2rem!important;}.stButton>button,button[kind="primary"]{min-height:46px!important;width:100%!important;}[data-testid="stHorizontalBlock"]{flex-direction:column!important;gap:.55rem!important;}[data-testid="stHorizontalBlock"]>[data-testid="column"]{width:100%!important;min-width:100%!important;max-width:100%!important;flex:1 1 100%!important;}div[data-baseweb="tab-list"]{overflow-x:auto!important;overflow-y:hidden!important;flex-wrap:nowrap!important;scrollbar-width:none!important;-webkit-overflow-scrolling:touch!important;}button[data-baseweb="tab"]{flex:0 0 auto!important;min-width:max-content!important;white-space:nowrap!important;padding:9px 11px!important;font-size:.78rem!important;background:#f4f5f8!important;border:1px solid #e1e4eb!important;}button[data-baseweb="tab"][aria-selected="true"]{color:#4c1d95!important;background:linear-gradient(135deg,#eee9ff,#e5deff)!important;border-color:#c9bfff!important;}section[data-testid="stSidebar"]{width:min(78vw,280px)!important;}section[data-testid="stSidebar"]>div:first-child{width:min(78vw,280px)!important;}.sidebar-ai-agent{font-size:.95rem!important;}.sidebar-max{padding:9px 10px 7px!important;}.sidebar-max-title{font-size:1.05rem!important;}.sidebar-max-text{font-size:.68rem!important;}[data-baseweb="select"],[data-baseweb="input"],[data-testid="stTextArea"]{font-size:16px!important;}[data-testid="stDataFrame"],[data-testid="stTable"]{width:100%!important;overflow-x:auto!important;}.stMarkdown,.stCaption{overflow-wrap:anywhere!important;}}
[data-testid="stSidebar"] [data-testid="stToggle"] label,[data-testid="stSidebar"] [data-baseweb="checkbox"] label{color:#b8ff00!important;text-shadow:0 0 7px rgba(184,255,0,.65)!important;}</style>""",
    unsafe_allow_html=True,
)

# MAX launcher: direct dialog call. Keep the dialog itself as the only Streamlit fragment.
if st.button("⚡ MAX", key="mobile_max_launcher", type="primary"):
    st.session_state["open_max"] = True
    render_max()

st.markdown('<p class="main-title">AI AGENT CONTENT MANAGER</p><div class="mobile-nav-hint">Разделы · листайте меню влево и вправо</div>', unsafe_allow_html=True)

base_tab_labels = [
    "⌂ Главная", "＋ Товар", "▦ Каталог", "✎ Тексты", "▶ Видео",
    "◷ План", "◉ Аналитика", "⚙ Настройки"
]
_is_platform_admin = platform_admin_enabled()
tab_labels = (["♛ АДМИН"] + base_tab_labels) if _is_platform_admin else base_tab_labels
_tabs = st.tabs(tab_labels)
if _is_platform_admin:
    tab_admin = _tabs[0]
    tab_dashboard, tab1, tab2, tab3, tab4, tab5, tab6, tab7 = _tabs[1:9]
else:
    tab_admin = None
    tab_dashboard, tab1, tab2, tab3, tab4, tab5, tab6, tab7 = _tabs[:8]

# ========== DASHBOARD ==========
if tab_admin is not None:
    with tab_admin:
        render_platform_admin()

with tab_dashboard:
    render_dashboard(
        load_products=load_products, load_plan=load_plan, load_leads=load_leads, load_orders=load_orders,
        crm_metrics=crm_metrics, order_metrics=order_metrics, conversion_metrics=conversion_metrics,
        workflow_metrics=workflow_metrics, low_stock_products=low_stock_products,
        growth_recommendations=growth_recommendations, max_product_title=max_product_title,
        attribution_loader=load_attribution, attribution_metrics=attribution_metrics,
        add_product=add_product, categories=CATEGORIES, can_write=can("write_data"),
    )

# ========== 1: СОЗДАТЬ КАРТОЧКУ ==========
with tab1:
    st.markdown('<div class="section-kicker">CONTENT STUDIO</div><div class="section-title">Создать товар</div><div class="section-subtitle">Загрузите фото, заполните данные и сразу получите готовую карточку.</div>', unsafe_allow_html=True)
    if not can("write_data"):
        st.info("Ваша роль доступна только для просмотра. Создание товаров доступно пользователям с правом записи.")
    up = st.file_uploader("📷 Фото товара", type=["jpg","jpeg","png","webp"])
    c1, c2, c3 = st.columns(3)
    with c1:
        name = st.text_input("Название *")
        brand = st.text_input("Бренд *")
        article = st.text_input("Артикул")
    with c2:
        sizes = st.text_input("Размеры")
        color = st.text_input("Цвет")
        category = st.selectbox("Категория", CATEGORIES)
    with c3:
        description = st.text_area("Описание", height=100)
        template = st.selectbox("🎨 Шаблон", list(TEMPLATES.keys()))
    specs = st.text_input("Характеристики через запятую")

    if can("write_data") and st.button("🎨 Создать карточку", type="primary"):
        if not name or not brand:
            st.error("Заполните: Название и Бренд")
        else:
            with st.spinner("Генерация..."):
                img = Image.open(up).convert("RGB") if up else None
                card = generate_card(img, name, brand, article, sizes, color, description, specs, category, template)
                card_buf = io.BytesIO()
                card.save(card_buf, format="PNG")
                card_b64 = base64.b64encode(card_buf.getvalue()).decode("ascii")

                original_b64 = None
                if up:
                    original_buf = io.BytesIO()
                    img.save(original_buf, format="JPEG", quality=95)
                    original_b64 = base64.b64encode(original_buf.getvalue()).decode("ascii")

                product_data = {
                    "name": name, "brand": brand, "article": article,
                    "sizes": sizes, "color": color, "description": description,
                    "specs": specs, "category": category,
                    "card_image": card_b64,
                    "date_added": str(datetime.date.today())
                }
                if original_b64:
                    product_data["original_image"] = original_b64

                add_product(product_data)
                st.success("✅ Карточка создана!")
                st.image(card, caption="Готово (1080×1350)", width=400)
                buf = io.BytesIO()
                card.save(buf, format="PNG")
                st.session_state["last_card_bytes"] = buf.getvalue()
                st.session_state["last_card_name"] = name
                st.download_button("⬇️ Скачать карточку", buf.getvalue(),
                    file_name=f"{name.replace(' ','_')}_card.png", mime="image/png")


# ==================== МАССОВЫЙ ИМПОРТ ====================
BULK_FIELDS = ["name", "brand", "article", "sizes", "color", "category", "description", "specs"]

def _bulk_value(row, *keys):
    normalized = {str(k).strip().lower().replace(" ", "_"): v for k, v in row.items()}
    aliases = {
        "name": ["name", "название", "товар"],
        "brand": ["brand", "бренд"],
        "article": ["article", "артикул"],
        "sizes": ["sizes", "размеры", "размер"],
        "color": ["color", "цвет"],
        "category": ["category", "категория"],
        "description": ["description", "описание"],
        "specs": ["specs", "характеристики"],
    }
    for key in keys:
        for alias in aliases.get(key, [key]):
            value = normalized.get(alias)
            if value is not None and str(value).strip():
                return str(value).strip()
    return ""

def parse_bulk_file(uploaded_file):
    raw = uploaded_file.getvalue()
    name = (uploaded_file.name or "").lower()

    if name.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise RuntimeError("Для XLSX нужен openpyxl. CSV можно загружать без дополнительных зависимостей.")
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(x or "").strip() for x in rows[0]]
        return [dict(zip(headers, row)) for row in rows[1:] if any(x not in (None, "") for x in row)]
    text = raw.decode("utf-8-sig")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\\t,")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";"
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    return [dict(row) for row in reader if any(str(v or "").strip() for v in row.values())]

def bulk_import_products(uploaded_file, update_existing=False):
    rows = parse_bulk_file(uploaded_file)
    products = load_products()
    existing = {
        str(p.get("article", "")).strip().lower(): i
        for i, p in enumerate(products)
        if str(p.get("article", "")).strip()
    }
    added = updated = skipped = 0
    errors = []

    if not can("write_data"):
        raise PermissionError("У вашей роли нет прав на импорт товаров.")
    for line_no, row in enumerate(rows, start=2):
        name = _bulk_value(row, "name")
        brand = _bulk_value(row, "brand")
        article = _bulk_value(row, "article")
        if not name:
            skipped += 1
            errors.append(f"Строка {line_no}: нет названия.")
            continue

        category = _bulk_value(row, "category") or "Другое"
        if category not in CATEGORIES:
            category = "Другое"

        product = {
            "name": name,
            "brand": brand,
            "article": article,
            "sizes": _bulk_value(row, "sizes"),
            "color": _bulk_value(row, "color"),
            "description": _bulk_value(row, "description"),
            "specs": _bulk_value(row, "specs"),
            "category": category,
            "date_added": str(datetime.date.today()),
        }

        key = article.lower()
        if update_existing and key and key in existing:
            products[existing[key]].update(product)
            updated += 1
        else:
            if not feature_allowed("products", len(products)):
                raise ValueError(f"Лимит каталога тарифа {tenant_plan().upper()} достигнут.")
            products.append(product)
            if key:
                existing[key] = len(products) - 1
            added += 1

    save_products(products)
    return added, updated, skipped, errors


# ========== 2: КАТАЛОГ ==========
with tab2:
    render_catalog(
        categories=CATEGORIES,
        category_emoji=CATEGORY_EMOJI,
        data_load_page=data_load_page,
        update_product=update_product,
        delete_product=delete_product,
        bulk_import_products=bulk_import_products,
        data_conflict_error=DataConflictError,
        can_write=can("write_data"),
    )

# ========== 3: ТЕКСТЫ С ПУБЛИКАЦИЕЙ ==========
with tab3:
    st.header("Тексты для соцсетей")
    products = load_products()
    if not products:
        st.info("Сначала создайте товар.")
    else:
        names = [f"{p.get('brand','')} {p.get('name','')} ({p.get('article','')})" for p in products]

        idx = st.selectbox(
            "Товар",
            range(len(names)),
            format_func=lambda x: names[x],
            key="text_product"
        )

        tone = st.radio(
            "Тональность",
            TONES,
            horizontal=True,
            key="text_tone"
        )

        st.caption(f"Выбрано: **{names[idx]}** · Тональность: **{tone}**")

        if st.button(
            "✨ Сгенерировать тексты",
            type="primary",
            use_container_width=True,
            key="generate_texts_btn"
        ):
            p = products[idx]

            # Каждый запуск заново создаёт тексты именно для выбранной тональности.
            st.session_state["tg_text"] = gen_telegram(p, tone)
            st.session_state["insta_text"] = gen_instagram(p, tone)
            st.session_state["vk_text"] = gen_vk(p, tone)
            st.session_state["current_product"] = p
            st.session_state["current_tone"] = tone
            st.session_state["text_generation_id"] = st.session_state.get("text_generation_id", 0) + 1

        if "tg_text" in st.session_state:
            p = st.session_state["current_product"]
            insta = st.session_state["insta_text"]
            tg = st.session_state["tg_text"]
            vk = st.session_state["vk_text"]
            generated_tone = st.session_state.get("current_tone", "")

            st.success(
                f"Тексты созданы: {p.get('brand','')} {p.get('name','')} · «{generated_tone}»"
            )

            st.caption(
                "Показаны тексты для последней нажатой кнопки «Сгенерировать тексты»."
            )

            st.markdown("---")
            st.subheader("📸 Instagram")
            st.text_area(
                "Текст поста",
                insta,
                height=280,
                key=f"i_out_{st.session_state.get('text_generation_id', 0)}"
            )
            st.download_button(
                "⬇️ Скачать текст",
                insta,
                file_name=f"{p.get('name','item')}_insta.txt",
                mime="text/plain",
                key=f"dl_i_{st.session_state.get('text_generation_id', 0)}"
            )
            st.caption("ℹ️ Скопируйте текст и загрузите карточку в Instagram вручную.")

            st.markdown("---")
            st.subheader("✈️ Telegram")
            st.text_area(
                "Текст поста",
                tg,
                height=280,
                key=f"t_out_{st.session_state.get('text_generation_id', 0)}"
            )

            card_to_send = None
            if "last_card_bytes" in st.session_state:
                card_to_send = st.session_state["last_card_bytes"]
            elif p.get("card_image"):
                try:
                    card_to_send = base64.b64decode(p["card_image"])
                except Exception:
                    card_to_send = None

            col_a, col_b = st.columns(2)
            with col_a:
                if can("write_data") and st.button(
                    "🚀 Опубликовать в Telegram",
                    type="primary",
                    key=f"pub_tg_text_{st.session_state.get('text_generation_id', 0)}"
                ):
                    if not card_to_send:
                        st.error("Нет карточки. Создайте её во вкладке «📸 Создать» или загрузите файл.")
                    else:
                        ok, msg = publish_to_telegram(card_to_send, tg)
                        (st.success if ok else st.error)(msg)
            with col_b:
                st.download_button(
                    "⬇️ Скачать текст",
                    tg,
                    file_name=f"{p.get('name','item')}_tg.txt",
                    mime="text/plain",
                    key=f"dl_tg_{st.session_state.get('text_generation_id', 0)}"
                )

            st.markdown("---")
            st.subheader("🅥 ВКонтакте")
            st.text_area(
                "Текст поста",
                vk,
                height=200,
                key=f"v_out_{st.session_state.get('text_generation_id', 0)}"
            )
            st.download_button(
                "⬇️ Скачать текст",
                vk,
                file_name=f"{p.get('name','item')}_vk.txt",
                mime="text/plain",
                key=f"dl_vk_{st.session_state.get('text_generation_id', 0)}"
            )

# ========== 4: REELS & STORIES ==========
with tab4:
    st.markdown('<div class="section-kicker">SHORT VIDEO</div><div class="section-title">Reels & Stories</div><div class="section-subtitle">Идея → сценарий → видео → публикация.</div>', unsafe_allow_html=True)
    products = load_products()
    if not products:
        st.info("Сначала создайте товар.")
    else:
        names = [f"{p.get('brand','')} {p.get('name','')}" for p in products]
        idx = st.selectbox("Товар", range(len(names)), format_func=lambda x: names[x], key="r_sel")

        if st.button("🎬 Сгенерировать", type="primary"):
            p = products[idx]
            n, b, d = p.get('name',''), p.get('brand',''), p.get('description','')

            ideas = f"""💡 ИДЕИ ДЛЯ REELS:

1. 📦 Распаковка — показать товар крупным планом
2. 👕 Обзор — надеть и показать в движении
3. 🎨 Стилизация — 3 образа с товаром
4. ⚖️ Сравнение — до/после
5. ❓ Q&A — ответы на вопросы о {n}"""

            script = f"""🎬 СЦЕНАРИЙ REELS (30 сек):

[0-5 сек] Открывашка
📹 Общий план товара
💬 "Смотрите, какая новинка!"
🎙️ "Встречайте {n} от {b}!"

[5-15 сек] Детали
📹 Крупный план
💬 "Качество, которое видно"
🎙️ "{d}"

[15-25 сек] Примерка
📹 Человек с товаром
💬 "Идеально подходит для..."
🎙️ "Размеры: {p.get('sizes','уточняйте')}. Цвет: {p.get('color','уточняйте')}"

[25-30 сек] CTA
📹 Логотип AI Agent Content Manager
💬 "Заказывайте!"
🎙️ "Ссылка в шапке профиля!"

🎵 Энергичный спортивный трек"""

            stories = f"""📱 ИДЕИ STORIES:

1. Опрос: "Какой цвет лучше?"
2. Слайдер: "Оцените от 1 до 10"
3. Q&A: "Задайте вопрос о {n}"
4. Обратный отсчёт до старта
5. Ссылка на товар"""

            st.subheader("💡 Идеи")
            st.text_area("", ideas, height=200, key="id_out")
            st.download_button("⬇️ Идеи", ideas, file_name=f"{n}_ideas.txt")

            st.subheader("🎬 Сценарий")
            st.text_area("", script, height=350, key="sc_out")
            st.download_button("⬇️ Сценарий", script, file_name=f"{n}_script.txt")

            st.subheader("📱 Stories")
            st.text_area("", stories, height=150, key="st_out")
            st.download_button("⬇️ Stories", stories, file_name=f"{n}_stories.txt")

        st.markdown("---")

        st.subheader("🎞️ Видео-рилс из карточки")

        # По умолчанию Reels использует готовую карточку выбранного товара.
        reel_product = products[st.session_state.get("r_sel", 0)]
        stored_card = reel_product.get("card_image")

        # Совместимость со старыми товарами.
        if not stored_card and st.session_state.get("last_card_name") == reel_product.get("name"):
            stored_card = base64.b64encode(
                st.session_state.get("last_card_bytes", b"")
            ).decode("ascii")

        replace_reel_photo = st.checkbox(
            "🔄 Заменить фото для этого Reels",
            value=False,
            key="replace_reel_photo"
        )

        reel_file = None
        if replace_reel_photo:
            reel_file = st.file_uploader(
                "Выберите другое изображение",
                type=["jpg", "jpeg", "png", "webp"],
                key="reel_card_replace"
            )

        if stored_card and not replace_reel_photo:
            st.info(
                f"📎 Используется карточка товара: "
                f"**{reel_product.get('brand','')} {reel_product.get('name','')}**"
            )
        elif not stored_card and not replace_reel_photo:
            st.info(
                "ℹ️ У этого товара нет сохранённой карточки. "
                "При создании Reels карточка будет создана автоматически "
                "из данных товара и сохранена в каталоге."
            )

        if st.button("🎬 Создать видео", key="make_reel_btn"):
            import tempfile, os
            from reels import make_reel

            if replace_reel_photo:
                if not reel_file:
                    st.error("Выбери изображение для замены.")
                    st.stop()
                source_bytes = reel_file.getvalue()
            else:
                if not stored_card:
                    source_image = None
                    original_b64 = reel_product.get("original_image")
                    if original_b64:
                        try:
                            source_image = Image.open(
                                io.BytesIO(base64.b64decode(original_b64))
                            ).convert("RGB")
                        except Exception:
                            source_image = None

                    if source_image is None:
                        st.error(
                            "❌ Нельзя создать Reels: у товара нет исходного фото. "
                            "Открой «📦 Каталог → ✏️ Редактировать», загрузи фото товара и сохрани."
                        )
                        st.stop()

                    auto_card = generate_card(
                        source_image,
                        reel_product.get("name", ""),
                        reel_product.get("brand", ""),
                        reel_product.get("article", ""),
                        reel_product.get("sizes", ""),
                        reel_product.get("color", ""),
                        reel_product.get("description", ""),
                        reel_product.get("specs", ""),
                        reel_product.get("category", "Другое"),
                        "Спортивный",
                        get_logo(),
                    )
                    card_buf = io.BytesIO()
                    auto_card.save(card_buf, format="PNG")
                    source_bytes = card_buf.getvalue()

                    updated_product = dict(reel_product)
                    updated_product["card_image"] = base64.b64encode(source_bytes).decode("ascii")
                    update_product(reel_product.get("_saas_record_id") or st.session_state.get("r_sel", 0), updated_product)
                    reel_product = updated_product
                    stored_card = updated_product["card_image"]

                    st.success(
                        "✅ Карточка автоматически создана с исходным фото товара и сохранена."
                    )
                else:
                    source_bytes = base64.b64decode(stored_card)

            with tempfile.TemporaryDirectory() as tmp:
                src = os.path.join(tmp, "card.png")
                with open(src, "wb") as f:
                    f.write(source_bytes)

                out = os.path.join(tmp, "reel.mp4")
                with st.spinner("Рендерю видео, подожди..."):
                    make_reel(src, out, duration=8)

                with open(out, "rb") as f:
                    video_bytes = f.read()

            # ВАЖНО: сохраняем MP4 в session_state. Streamlit перезапускает
            # скрипт при нажатии кнопки публикации, поэтому локальная переменная
            # video_bytes иначе теряется.
            st.session_state["reel_video_bytes"] = video_bytes
            st.session_state["reel_caption"] = gen_instagram(reel_product, "Продающий")
            st.session_state["reel_product_name"] = (
                f"{reel_product.get('brand','')} {reel_product.get('name','')}".strip()
            )
            st.success("✅ Reels создан и сохранён. Теперь можно публиковать.")

        # После создания MP4 этот блок остаётся доступным на следующих rerun.
        video_bytes = st.session_state.get("reel_video_bytes")
        reel_caption = st.session_state.get("reel_caption", "")
        reel_product_name = st.session_state.get("reel_product_name", "")

        if video_bytes:
            st.video(video_bytes, width=260)

            st.download_button(
                "⬇️ Скачать рилс",
                video_bytes,
                "reel.mp4",
                "video/mp4",
                key="dl_reel"
            )

            st.markdown("### 📤 Опубликовать Reels")
            if reel_product_name:
                st.caption(f"Товар: {reel_product_name}")

            pub1, pub2, pub3 = st.columns(3)

            with pub1:
                if can("write_data") and st.button("✈️ В Telegram", key="publish_reel_tg"):
                    with st.spinner("Отправляю Reels в Telegram..."):
                        ok, msg = publish_reel_to_telegram(video_bytes, reel_caption)
                    (st.success if ok else st.error)(msg)

            with pub2:
                if can("write_data") and st.button("🅥 В VK", key="publish_reel_vk"):
                    with st.spinner("Отправляю Reels в VK..."):
                        ok, msg = publish_reel_to_vk(video_bytes, reel_caption)
                    (st.success if ok else st.error)(msg)

            with pub3:
                st.button(
                    "📸 В Instagram",
                    key="publish_reel_instagram",
                    disabled=True,
                    help="Для автоматической публикации Instagram требует публичный URL видео и подключение Meta API."
                )

            st.caption(
                "Telegram и VK публикуются непосредственно из приложения. "
                "Instagram подключим после добавления публичного хранилища для MP4 и Meta API."
            )

        st.markdown("---")


# ========== 5: КОНТЕНТ-ПЛАН ==========
with tab5:
    st.markdown('<div class="section-kicker">CONTENT WORKFLOW</div><div class="section-title">Контент-центр</div><div class="section-subtitle">Единый процесс: идея → работа → проверка → готово → публикация.</div>', unsafe_allow_html=True)
    products = load_products()
    plan = load_plan()

    wm = workflow_metrics(plan)
    m1, m2, m3, m4, m5 = st.columns(5)
    m1.metric("Всего", wm["total"])
    m2.metric("В работе", wm["counts"].get("В работе", 0))
    m3.metric("На проверке", wm["counts"].get("На проверке", 0))
    m4.metric("Готово", wm["counts"].get("Готово", 0))
    m5.metric("Опубликовано", wm["counts"].get("Опубликовано", 0))
    if wm["overdue"]:
        st.warning(f"Просрочено: {wm['overdue']}")

    planner_tab, workflow_tab, adapt_tab, report_tab = st.tabs(["📅 План", "🔄 Workflow", "📣 Адаптация", "📊 Отчёт"])

    with planner_tab:
        if not products:
            st.info("Сначала создайте товар.")
        else:
            with st.form("plan_f", clear_on_submit=True):
                c1, c2, c3 = st.columns(3)
                with c1:
                    pd = st.date_input("Дата")
                    pl = st.selectbox("Платформа", PLATFORMS)
                with c2:
                    pn = [f"{p.get('brand','')} {p.get('name','')}".strip() for p in products]
                    sp = st.selectbox("Товар", pn)
                    ct = st.selectbox("Тип", CONTENT_TYPES)
                with c3:
                    sts = st.selectbox("Статус", STATUSES)
                    pr = st.selectbox("Приоритет", PRIORITIES)
                idea = st.text_area("Идея / текст")
                if can("write_data") and st.form_submit_button("📌 Добавить в workflow"):
                    add_plan(ensure_workflow({"date": str(pd), "platform": pl, "product": sp,
                                              "type": ct, "idea": idea, "status": sts, "priority": pr}))
                    st.success("Материал добавлен.")
                    st.rerun()

            if plan:
                st.subheader("📅 Календарь по неделям")
                weeks = {}
                for i, raw_item in enumerate(plan):
                    item = ensure_workflow(raw_item)
                    try:
                        d = datetime.date.fromisoformat(str(item.get("date", "")))
                        wk = d.isocalendar()[1]
                    except Exception:
                        wk = 0
                    weeks.setdefault(wk, []).append((i, item))
                for wk in sorted(weeks.keys()):
                    with st.expander(f"Неделя {wk} · {len(weeks[wk])} материалов"):
                        for i, item in weeks[wk]:
                            st.write(f"**{item.get('date','')}** · {item.get('platform','')} · {item.get('type','')} · **{item.get('status','Идея')}**")
                            st.caption(f"{item.get('product','')} — {item.get('idea','')}")

            st.markdown("---")
            st.subheader("Управление материалами")
            for i, raw_item in enumerate(reversed(plan)):
                ri = len(plan) - 1 - i
                item = ensure_workflow(raw_item)
                with st.expander(f"{item.get('date','')} · {item.get('product','')} · {item.get('platform','')}"):
                    st.write(f"**Идея:** {item.get('idea','')}")
                    cs = item.get("status", "Идея")
                    ns = st.selectbox("Этап", STATUSES, index=STATUSES.index(cs) if cs in STATUSES else 0, key=f"st_{ri}", disabled=not can("write_data"))
                    if can("write_data") and ns != cs:
                        updated = change_status(item, ns)
                        update_plan(item.get("_saas_record_id") or ri, updated)
                        st.rerun()
                    st.caption(f"Приоритет: {item.get('priority','Обычный')}")
                    if item.get("history"):
                        st.caption("История изменений")
                        for h in item["history"][-5:]:
                            st.write(f"{h.get('time','')} · {h.get('from','')} → {h.get('to','')} · {h.get('actor','manager')}")
                    if can("write_data") and st.button("🗑️ Удалить", key=f"dp_{ri}"):
                        delete_plan(item.get("_saas_record_id") or ri)
                        st.rerun()

    with workflow_tab:
        st.subheader("🔄 Очередь на проверку")
        review_items = [(i, ensure_workflow(x)) for i, x in enumerate(plan) if ensure_workflow(x).get("status") == "На проверке"]
        if not review_items:
            st.info("Материалов на проверке пока нет.")
        else:
            for i, item in review_items:
                st.markdown(f"**{item.get('product','')} · {item.get('platform','')} · {item.get('type','')}**")
                st.write(item.get("idea", ""))
                a, b = st.columns(2)
                with a:
                    if can("write_data") and st.button("✅ Утвердить", key=f"approve_{i}"):
                        update_plan(item.get("_saas_record_id") or i, change_status(item, "Готово"))
                        st.rerun()
                with b:
                    if can("write_data") and st.button("↩️ Вернуть в работу", key=f"return_{i}"):
                        update_plan(item.get("_saas_record_id") or i, change_status(item, "В работе"))
                        st.rerun()
        st.markdown("---")
        st.subheader("🤖 AI-рекомендации")
        for rec in recommendations(products, plan):
            st.write(f"• {rec}")

    with adapt_tab:
        st.subheader("📣 Один материал → три канала")
        if not products:
            st.info("Сначала добавьте товары.")
        else:
            names = [f"{p.get('brand','')} {p.get('name','')}".strip() for p in products]
            sel = st.selectbox("Товар", names, key="content_adapt_product")
            p = products[names.index(sel)]
            base_item = {"product": sel, "type": "Пост", "status": "Идея"}
            adapted = adapt_content(base_item, p, {"Instagram": gen_instagram, "Telegram": gen_telegram, "VK": gen_vk})
            for channel in ["Instagram", "Telegram", "VK"]:
                st.markdown(f"**{channel}**")
                st.text_area(channel, adapted.get(channel, ""), height=150, key=f"adapt_{channel}")
            st.caption("Цены в карточки и тексты автоматически не добавляются.")

    with report_tab:
        st.subheader("📊 Отчёт и история")
        for line in report_lines(products, plan, load_leads(), load_orders()):
            st.write(line)
        report_text = "\n".join(report_lines(products, plan, load_leads(), load_orders()))
        st.download_button("⬇️ Скачать отчёт TXT", report_text, file_name=f"arsenal_content_report_{datetime.date.today()}.txt")
        csv_buf = io.StringIO()
        writer = csv.writer(csv_buf)
        writer.writerow(["Дата","Платформа","Товар","Тип","Статус","Приоритет","Идея","История"])
        for item in plan:
            writer.writerow([item.get("date",""), item.get("platform",""), item.get("product",""),
                             item.get("type",""), item.get("status","Идея"), item.get("priority","Обычный"),
                             item.get("idea",""), json.dumps(item.get("history", []), ensure_ascii=False)])
        st.download_button("⬇️ Экспорт истории CSV", csv_buf.getvalue(),
                           file_name=f"arsenal_content_history_{datetime.date.today()}.csv", mime="text/csv")

# ========== 6: СТАТИСТИКА ==========
with tab6:
    st.markdown('<div class="section-kicker">GROWTH ENGINE</div><div class="section-title">Контент · Продажи · Аналитика · AI</div><div class="section-subtitle">Единый контур: товар → контент → публикация → заявка → заказ → выручка → следующая рекомендация.</div>', unsafe_allow_html=True)
    products = load_products()
    plan = load_plan()
    leads = load_leads()
    orders = load_orders()
    growth = ai_summary(products, leads, orders, plan)
    f = growth["funnel"]

    c1, c2, c3, c4, c5 = st.columns(5)
    c1.metric("Контент", f["content"])
    c2.metric("Опубликовано", f["published"])
    c3.metric("Заявки", f["leads"])
    c4.metric("Заказы", f["orders"])
    c5.metric("Выручка", f"{f['revenue']:,.0f} ₽".replace(",", " "))

    st.markdown("---")
    a, b = st.columns(2)
    with a:
        st.subheader("Воронка")
        st.write(f"Контент: **{f['content']}**")
        st.write(f"Опубликовано: **{f['published']}**")
        st.write(f"Заявки: **{f['leads']}**")
        st.write(f"Заказы: **{f['orders']}**")
        st.write(f"Завершённые заказы: **{f['completed_orders']}**")
        st.write(f"Конверсия заявка → заказ: **{f['lead_to_order']:.1f}%**")
    with b:
        st.subheader("🤖 AI-рекомендации")
        for rec in growth["recommendations"]:
            st.write(f"• {rec}")

    st.markdown("---")
    st.subheader("🔗 Реальная атрибуция: публикация → заявка → заказ → выручка")
    attribution_rows = attribution_performance(plan, leads, orders)
    explicit = [x for x in attribution_rows if x["leads"] or x["orders"] or x["revenue"]]
    if explicit:
        st.caption("Показываются только явные связи по content_id. Это фактическая атрибуция события, а не вывод по совпадению названий.")
        for row in explicit[:20]:
            st.write(f"**{row['date']} · {row['channel']} · {row['format']}** · {row['product']} — заявки {row['leads']} · заказы {row['orders']} · {row['revenue']:,.0f} ₽".replace(",", " "))
    else:
        st.info("Явных связей пока нет. При создании заказа можно выбрать опубликованный материал — связь сохранится автоматически.")

    st.subheader("📦 Товары: связь контента и продаж")
    st.caption("Связь определяется по названию/артикулу товара в заявке, заказе и контент-плане; это атрибуция по совпадению, а не доказательство причинности.")
    for row in growth["products"][:20]:
        st.write(f"**{row['product']}** · контент {row['content']} · заявки {row['leads']} · заказы {row['orders']} · {row['revenue']:,.0f} ₽ · остаток {row['stock']} шт.".replace(",", " "))

    st.markdown("---")
    st.subheader("🚀 Следующий контент")
    for item in growth["next_content"]:
        st.write(f"• **{item['priority']}** · {item['type']} · {item['product']} — {item['idea']}")

    st.markdown("---")
    if products:
        st.subheader("📦 Категории")
        cc = {}
        for p in products:
            c = p.get('category','Другое')
            cc[c] = cc.get(c,0) + 1
        for c, n in sorted(cc.items(), key=lambda x: -x[1]):
            st.write(f"{CATEGORY_EMOJI.get(c,'📦')} **{c}:** {n}")

    if plan:
        st.subheader("📅 Статусы")
        sc = {}
        for item in plan:
            s = item.get('status','Идея')
            sc[s] = sc.get(s,0) + 1
        for s, n in sorted(sc.items(), key=lambda x: -x[1]):
            st.write(f"**{s}:** {n}")

# ==================== НАСТРОЙКИ МАГАЗИНА ====================
SETTINGS_FILE = Path("settings.json")
DEFAULT_SETTINGS = {
    "onboarding_complete": True,
    "store_name": "AI Agent Content Manager",
    "city": "Краснодар",
    "delivery": "По России",
    "telegram": "",
    "vk": "",
    "instagram": "",
    "order_contact": "",
    "phone": "",
    "manager_name": "",
    "pickup_address": "",
    "payment_methods": "Наличные, перевод, карта",
    "return_policy": "Возврат и обмен — по действующим правилам магазина.",
    "delivery_methods": "Самовывоз, доставка по Краснодару, отправка по России",
    "delivery_terms": "",
    "content_signature": "AI Agent Content Manager — спортивная одежда, обувь и экипировка",
    "cta": "Напишите нам — поможем выбрать товар и оформить заказ.",
    "hashtags": "#ai_agent_content_manager #спорт #экипировка",
    "ai_seller_instructions": "Отвечай дружелюбно и по делу. Уточняй товар, размер и город. Не придумывай наличие, цену или характеристики.",
    "notification_contact": "",
    "card_template": "Dark Premium",
    "show_price_on_cards": "Нет",
    "main_sport": "Футбол",
    "ai_tone": "Дружелюбный и профессиональный",
    "ai_required_questions": "Товар, размер, город, способ получения",
    "ai_no_stock_reply": "Если товара нет, предложи похожие варианты из каталога.",
    "ai_escalation_reply": "Если вопрос нельзя решить по данным магазина, предложи связаться с менеджером.",
    "faq": [],
}

def load_settings():
    rows = data_load("settings", [])
    data = rows[0] if isinstance(rows, list) and rows and isinstance(rows[0], dict) else {}
    result = DEFAULT_SETTINGS.copy()
    result.update({k: data.get(k, DEFAULT_SETTINGS[k]) for k in DEFAULT_SETTINGS})
    return result

def save_settings(data):
    payload = {k: data.get(k, "") for k in DEFAULT_SETTINGS}
    existing = data_load("settings", [])
    if existing and isinstance(existing[0], dict) and existing[0].get("_saas_record_id"):
        payload["_saas_record_id"] = existing[0]["_saas_record_id"]
    data_save("settings", [payload])



# ========== 7: НАСТРОЙКИ ==========
with tab7:
    if can("settings"):
        render_settings(
            load_settings=load_settings, save_settings=save_settings, templates=TEMPLATES,
            get_logo=get_logo, save_logo=save_logo,
        )
    else:
        st.info("Настройки магазина доступны только владельцу и администратору.")
